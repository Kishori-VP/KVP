"""Render report tables to PDF (fpdf2, with Marathi text shaping), Excel (openpyxl), CSV and Word (python-docx).

A report is a dict:
{
  "title": str, "subtitle": str | None,
  "period": {"start": date, "end": date, "label": str} | None,
  "summary": [{"label": str, "value": any}],
  "sections": [{"title": str, "columns": [{"key": str, "label": str}], "rows": [dict]}],
  "notes": [str],
}
"""

import csv
import io
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Literal

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.utils.dates import local_tz

FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF = "application/pdf"
CSV = "text/csv; charset=utf-8"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ZIP = "application/zip"
ExportFormat = Literal["pdf", "xlsx", "csv", "docx"]
FORMAT_LABELS = {"pdf": "PDF", "xlsx": "Excel", "csv": "CSV", "docx": "Word"}
BRAND = "7B2D8E"


def _fmt(value) -> str | int | float:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.astimezone(local_tz()).strftime("%d-%m-%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (list, tuple)):
        return ", ".join(str(_fmt(v)) for v in value)
    if isinstance(value, dict):
        return ", ".join(f"{k}: {_fmt(v)}" for k, v in value.items())
    if isinstance(value, float):
        return round(value, 2)
    if isinstance(value, str) and _CODE.fullmatch(value):
        return value.replace("_", " ").capitalize()  # status codes like no_recent_activity
    return value


_CODE = re.compile(r"[a-z]+(_[a-z]+)+")


def _generated_line() -> str:
    return f"Generated {datetime.now(local_tz()):%d-%m-%Y %H:%M} · {settings.app_name}"


# ---------------------------------------------------------------- Excel

def to_xlsx(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    title_font = Font(bold=True, size=14, color=BRAND)
    head_fill = PatternFill("solid", fgColor=BRAND)
    head_font = Font(bold=True, color="FFFFFF")

    ws.append([report["title"]])
    ws["A1"].font = title_font
    if report.get("subtitle"):
        ws.append([report["subtitle"]])
    if report.get("period"):
        ws.append([f"Period: {report['period']['label']}"])
    ws.append([_generated_line()])
    ws.append([])
    if report.get("summary"):
        ws.append(["Indicator", "Value"])
        for cell in ws[ws.max_row]:
            cell.fill, cell.font = head_fill, head_font
        for item in report["summary"]:
            ws.append([item["label"], _fmt(item["value"])])
        ws.append([])
    for note in report.get("notes", []):
        ws.append([note])
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 22

    used = set()
    for section in report.get("sections", []):
        name = (section["title"][:28] or "Data").replace("/", "-").replace(":", "")
        base, i = name, 2
        while name in used:
            name = f"{base[:25]}-{i}"
            i += 1
        used.add(name)
        sh = wb.create_sheet(name)
        cols = section["columns"]
        sh.append([col["label"] for col in cols])
        for cell in sh[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in section["rows"]:
            sh.append([_fmt(row.get(col["key"])) for col in cols])
        if section.get("totals"):
            sh.append([_fmt(section["totals"].get(col["key"], "Total" if i == 0 else "")) for i, col in enumerate(cols)])
            for cell in sh[sh.max_row]:
                cell.font = Font(bold=True)
        for idx, col in enumerate(cols, start=1):
            longest = max([len(str(col["label"]))] + [len(str(_fmt(r.get(col["key"])))) for r in section["rows"][:500]])
            sh.column_dimensions[get_column_letter(idx)].width = min(max(10, longest + 2), 60)
        sh.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

class _PDF(FPDF):
    report_title = ""

    def header(self):
        self.set_font("NotoSans", "B", 9)
        self.set_text_color(123, 45, 142)
        self.cell(0, 6, f"{settings.app_name} — {self.report_title}", new_x="LMARGIN", new_y="NEXT")
        self.set_text_color(0, 0, 0)
        self.ln(1)

    def footer(self):
        self.set_y(-12)
        self.set_font("NotoSans", "", 7)
        self.set_text_color(120, 120, 120)
        self.cell(0, 8, f"{_generated_line()}   ·   Page {self.page_no()}/{{nb}}", align="C")


def _new_pdf(title: str, landscape: bool) -> _PDF:
    pdf = _PDF(orientation="L" if landscape else "P", unit="mm", format="A4")
    pdf.report_title = title
    pdf.add_font("NotoSans", "", str(FONT_DIR / "NotoSans-Regular.ttf"))
    pdf.add_font("NotoSans", "B", str(FONT_DIR / "NotoSans-Bold.ttf"))
    pdf.add_font("NotoDeva", "", str(FONT_DIR / "NotoSansDevanagari-Regular.ttf"))
    pdf.add_font("NotoDeva", "B", str(FONT_DIR / "NotoSansDevanagari-Bold.ttf"))
    pdf.set_fallback_fonts(["NotoDeva"])
    pdf.set_text_shaping(True)
    pdf.set_auto_page_break(True, margin=15)
    pdf.alias_nb_pages()
    pdf.add_page()
    return pdf


def _pdf_table(pdf: _PDF, columns: list[dict], rows: list[dict], totals: dict | None = None):
    pdf.set_font("NotoSans", "", 7.5 if len(columns) > 7 else 8.5)
    head = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=(123, 45, 142))
    # Size columns to their content (clamped), instead of equal widths.
    weights = []
    for col in columns:
        longest = max([len(str(col["label"]))] + [len(str(_fmt(r.get(col["key"])))) for r in rows[:300]])
        weights.append(min(max(longest, 7), 40))
    with pdf.table(headings_style=head, line_height=5, text_align="LEFT", padding=1, col_widths=tuple(weights),
                   cell_fill_color=(246, 240, 248), cell_fill_mode="ROWS") as table:
        hr = table.row()
        for col in columns:
            hr.cell(str(col["label"]))
        for r in rows:
            tr = table.row()
            for col in columns:
                tr.cell(str(_fmt(r.get(col["key"]))))
        if totals:
            tr = table.row()
            for i, col in enumerate(columns):
                tr.cell(str(_fmt(totals.get(col["key"], "Total" if i == 0 else ""))),
                        style=FontFace(emphasis="BOLD"))


def to_pdf(report: dict) -> bytes:
    widest = max([len(s["columns"]) for s in report.get("sections", [])] + [0])
    pdf = _new_pdf(report["title"], landscape=widest > 6)
    pdf.set_font("NotoSans", "B", 15)
    pdf.multi_cell(0, 8, report["title"], new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("NotoSans", "", 9.5)
    if report.get("subtitle"):
        pdf.multi_cell(0, 5, report["subtitle"], new_x="LMARGIN", new_y="NEXT")
    if report.get("period"):
        pdf.multi_cell(0, 5, f"Period: {report['period']['label']}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    if report.get("summary"):
        pdf.set_font("NotoSans", "B", 11)
        pdf.cell(0, 7, "Key indicators", new_x="LMARGIN", new_y="NEXT")
        _pdf_table(pdf, [{"key": "label", "label": "Indicator"}, {"key": "value", "label": "Value"}], report["summary"])
        pdf.ln(4)
    for section in report.get("sections", []):
        if not section["rows"]:
            continue
        pdf.set_font("NotoSans", "B", 11)
        pdf.cell(0, 7, section["title"], new_x="LMARGIN", new_y="NEXT")
        rows = section["rows"]
        truncated = len(rows) > 2000
        _pdf_table(pdf, section["columns"], rows[:2000], section.get("totals"))
        if truncated:
            pdf.set_font("NotoSans", "", 8)
            pdf.cell(0, 6, f"Showing first 2000 of {len(rows)} rows. Use the Excel export for the full list.",
                     new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)
    if report.get("notes"):
        pdf.set_font("NotoSans", "B", 10)
        pdf.cell(0, 6, "Notes", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("NotoSans", "", 8.5)
        for note in report["notes"]:
            pdf.multi_cell(0, 5, f"• {note}", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# ---------------------------------------------------------------- CSV

def to_csv(report: dict) -> bytes:
    """All sections in one sheet-like file, separated by blank lines. BOM so Excel reads Marathi correctly."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([report["title"]])
    if report.get("subtitle"):
        w.writerow([report["subtitle"]])
    if report.get("period"):
        w.writerow([f"Period: {report['period']['label']}"])
    w.writerow([_generated_line()])
    if report.get("summary"):
        w.writerow([])
        w.writerow(["Indicator", "Value"])
        for item in report["summary"]:
            w.writerow([item["label"], _fmt(item["value"])])
    for section in report.get("sections", []):
        cols = section["columns"]
        w.writerow([])
        w.writerow([section["title"]])
        w.writerow([col["label"] for col in cols])
        for row in section["rows"]:
            w.writerow([_fmt(row.get(col["key"])) for col in cols])
        if section.get("totals"):
            w.writerow([_fmt(section["totals"].get(col["key"], "Total" if i == 0 else "")) for i, col in enumerate(cols)])
    if report.get("notes"):
        w.writerow([])
        for note in report["notes"]:
            w.writerow([note])
    return buf.getvalue().encode("utf-8-sig")


# ---------------------------------------------------------------- Word

def _shade(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _docx_table(doc, columns: list[dict], rows: list[dict], totals: dict | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, col in zip(table.rows[0].cells, columns):
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(col["label"]))
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, BRAND)
    for r in rows:
        for cell, col in zip(table.add_row().cells, columns):
            cell.text = str(_fmt(r.get(col["key"])))
    if totals:
        for i, (cell, col) in enumerate(zip(table.add_row().cells, columns)):
            cell.text = ""
            cell.paragraphs[0].add_run(str(_fmt(totals.get(col["key"], "Total" if i == 0 else "")))).bold = True


def to_docx(report: dict) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Nirmala UI"  # ships with Windows and covers both Latin and Devanagari
    style.font.size = Pt(9)
    if max([len(s["columns"]) for s in report.get("sections", [])] + [0]) > 6:
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    doc.add_heading(report["title"], level=1)
    for line in (report.get("subtitle"), report.get("period") and f"Period: {report['period']['label']}", _generated_line()):
        if line:
            doc.add_paragraph(line)
    if report.get("summary"):
        doc.add_heading("Key indicators", level=2)
        _docx_table(doc, [{"key": "label", "label": "Indicator"}, {"key": "value", "label": "Value"}], report["summary"])
    for section in report.get("sections", []):
        if not section["rows"]:
            continue
        doc.add_heading(section["title"], level=2)
        rows = section["rows"]
        _docx_table(doc, section["columns"], rows[:2000], section.get("totals"))
        if len(rows) > 2000:
            doc.add_paragraph(f"Showing first 2000 of {len(rows)} rows. Use the Excel export for the full list.")
    if report.get("notes"):
        doc.add_heading("Notes", level=2)
        for note in report["notes"]:
            doc.add_paragraph(note, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


RENDERERS = {"pdf": (to_pdf, PDF), "xlsx": (to_xlsx, XLSX), "csv": (to_csv, CSV), "docx": (to_docx, DOCX)}


def render(report: dict, fmt: str) -> tuple[bytes, str, str]:
    fn, media_type = RENDERERS.get(fmt, RENDERERS["pdf"])
    return fn(report), media_type, fmt if fmt in RENDERERS else "pdf"


def render_many(report: dict, formats: list[str], stem: str,
                extra: list[tuple[str, bytes, str]] | None = None) -> tuple[bytes, str, str]:
    """One format gives that file; several formats, or any `extra` files (e.g. a template's attached
    cover letter), are bundled into a ZIP. Returns (content, media_type, filename)."""
    formats = list(dict.fromkeys(formats))  # de-duplicate, keep order
    if len(formats) == 1 and not extra:
        content, media_type, ext = render(report, formats[0])
        return content, media_type, f"{stem}.{ext}"
    buf = io.BytesIO()
    used: set[str] = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for fmt in formats:
            content, _, ext = render(report, fmt)
            used.add(f"{stem}.{ext}")
            zf.writestr(f"{stem}.{ext}", content)
        for name, content, _ in extra or []:
            base, dot, ext = name.rpartition(".")
            unique, i = name, 2
            while unique in used:
                unique = f"{base}-{i}{dot}{ext}" if dot else f"{name}-{i}"
                i += 1
            used.add(unique)
            zf.writestr(unique, content)
    return buf.getvalue(), ZIP, f"{stem}.zip"


# ---------------------------------------------------------------- event report

def event_report(ev: dict, fmt: str) -> tuple[bytes, str, str]:
    """ev is the serialized output of GET /events/{id}."""
    att = ev.get("attendance_summary", {})
    report_part = ev.get("report") or {}
    summary = [
        {"label": "Event", "value": ev["title"]},
        {"label": "Type", "value": ev["event_type"]},
        {"label": "Village", "value": (ev.get("village") or {}).get("name")},
        {"label": "Location", "value": ev.get("location_name")},
        {"label": "Date", "value": ev["local_date"]},
        {"label": "Time", "value": f"{ev.get('start_time') or ''} - {ev.get('end_time') or ''}"},
        {"label": "Status", "value": ev["display_status"]},
        {"label": "Responsible volunteer", "value": (ev.get("responsible_volunteer") or {}).get("name")},
        {"label": "Objective", "value": ev.get("objective")},
        {"label": "Expected participants", "value": att.get("expected")},
        {"label": "Girls present", "value": att.get("girls_present")},
        {"label": "Other participants", "value": att.get("other_participants")},
        {"label": "Total expenditure (₹)", "value": (ev.get("expenditure") or {}).get("total")},
        {"label": "Photos / media", "value": len(ev.get("media") or [])},
    ]
    sections = [
        {"title": "Post-event report", "columns": [{"key": "k", "label": "Item"}, {"key": "v", "label": "Details"}],
         "rows": [{"k": label, "v": report_part.get(key)} for key, label in (
             ("activities_conducted", "Activities conducted"), ("outcomes", "Outcomes"),
             ("observations", "Observations"), ("issues", "Issues"), ("follow_up_actions", "Follow-up actions"))]},
        {"title": "Attendance", "columns": [{"key": "girl_code", "label": "Girl ID"}, {"key": "name", "label": "Name"},
                                            {"key": "present", "label": "Present"}],
         "rows": ev.get("attendance") or []},
        {"title": "Pre-event checklist", "columns": [{"key": "item", "label": "Item"}, {"key": "done", "label": "Done"},
                                                     {"key": "notes", "label": "Notes"}],
         "rows": ev.get("pre_checklist") or []},
        {"title": "Post-event checklist", "columns": [{"key": "item", "label": "Item"}, {"key": "done", "label": "Done"},
                                                      {"key": "notes", "label": "Notes"}],
         "rows": ev.get("post_checklist") or []},
        {"title": "Expenditure", "columns": [{"key": "item", "label": "Item"}, {"key": "category", "label": "Category"},
                                             {"key": "quantity", "label": "Qty"}, {"key": "unit_cost", "label": "Unit cost"},
                                             {"key": "amount", "label": "Amount"}],
         "rows": (ev.get("expenditure") or {}).get("items") or [],
         "totals": {"amount": (ev.get("expenditure") or {}).get("total")}},
    ]
    report = {"title": f"Event Report: {ev['title']}", "subtitle": None, "period": None,
              "summary": summary, "sections": sections, "notes": []}
    return render(report, fmt)
